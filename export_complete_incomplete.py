"""Export live Kobo frequencies into Complete and Incomplete Excel sheets."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from dashboard.data import QUESTION_META, _column_or_grouped, load_union
from survey_tags import load_tag_map

ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "output"

AGE_ORDER = ["1", "2", "3", "4", "5", "6"]
GENDER_ORDER = ["1", "2"]
AGE_ELIGIBLE = {"2", "3", "4", "5"}

HEADER = PatternFill("solid", fgColor="121814")
TEAL = PatternFill("solid", fgColor="163832")
RED = PatternFill("solid", fgColor="B42318")
CREAM = PatternFill("solid", fgColor="F7F0E4")
WHITE = PatternFill("solid", fgColor="FFF8EE")
GOLD = Font(name="Calibri", color="C4A574", bold=True, size=11)
WHITE_FONT = Font(name="Calibri", color="FFF8EE", bold=True, size=11)
TITLE = Font(name="Calibri", bold=True, size=16, color="121814")
SUB = Font(name="Calibri", size=11, color="6F675C")
BODY = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="D4CBB8"),
    right=Side(style="thin", color="D4CBB8"),
    top=Side(style="thin", color="D4CBB8"),
    bottom=Side(style="thin", color="D4CBB8"),
)


def _has_value(series) -> bool:
    return series.notna() & series.astype(str).str.strip().ne("") & series.astype(str).ne("nan")


def classify(frame):
    q10 = _column_or_grouped(frame, "Q10_1")
    if q10.isna().all():
        q10 = _column_or_grouped(frame, "_Q10s1")
    complete = _has_value(q10)
    age_fail = frame["S1"].isin({"1", "6"})
    male = frame["S2"].eq("1")
    no_fragrance = frame["S3"].eq("2")
    qualified = frame["S1"].isin(AGE_ELIGIBLE) & frame["S2"].eq("2") & frame["S3"].eq("1")
    reason = []
    for i in frame.index:
        if bool(complete.loc[i]):
            reason.append("Complete interview")
        elif bool(age_fail.loc[i]):
            reason.append("Terminated · age out of range")
        elif bool(male.loc[i]):
            reason.append("Terminated · male")
        elif bool(no_fragrance.loc[i]):
            reason.append("Terminated · no fragrance in last 3 months")
        elif bool(qualified.loc[i]):
            reason.append("Did not finish after qualifying")
        else:
            reason.append("Did not finish")
    out = frame.copy()
    out["interview_status"] = ["Complete" if flag else "Incomplete" for flag in complete]
    out["status_reason"] = reason
    return out, complete


def freq(series, codes: list[str], labels: dict[str, str]) -> list[tuple[str, int, float]]:
    total = int(len(series))
    rows = []
    for code in codes:
        count = int(series.eq(code).sum())
        rows.append((labels[code], count, round((count / total) * 100, 1) if total else 0.0))
    missing = int(series.isna().sum() + (series.astype(str).isin({"", "nan", "None"})).sum())
    # na already counted in isna; avoid double-count of string nan
    missing = int(series.isna().sum())
    rows.append(("Not asked", missing, round((missing / total) * 100, 1) if total else 0.0))
    return rows


def write_cell(ws, row, col, value, fill=None, font=None, align=None, number=None):
    cell = ws.cell(row, col, value)
    cell.border = THIN
    cell.font = font or BODY
    cell.alignment = align or Alignment(horizontal="center", vertical="center", wrap_text=True)
    if fill:
        cell.fill = fill
    if number:
        cell.number_format = number
    return cell


def arm_block(group, arm: str):
    slice_ = group[group["arm"] == arm] if len(group) else group.iloc[0:0]
    return {
        "n": int(len(slice_)),
        "gender": freq(slice_["S2"] if len(slice_) else slice_.get("S2", slice_), GENDER_ORDER, QUESTION_META["S2"]["labels"]),
        "age": freq(slice_["S1"] if len(slice_) else slice_.get("S1", slice_), AGE_ORDER, QUESTION_META["S1"]["labels"]),
    }


def write_frequency_table(ws, start_row: int, title: str, rows: list[tuple[str, int, float]]) -> int:
    write_cell(ws, start_row, 1, title, fill=HEADER, font=WHITE_FONT, align=Alignment(horizontal="left", vertical="center"))
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
    write_cell(ws, start_row + 1, 1, "Answer", fill=CREAM, font=Font(name="Calibri", bold=True))
    write_cell(ws, start_row + 1, 2, "Count", fill=CREAM, font=Font(name="Calibri", bold=True))
    write_cell(ws, start_row + 1, 3, "%", fill=CREAM, font=Font(name="Calibri", bold=True))
    for i, (label, count, pct) in enumerate(rows, start=start_row + 2):
        write_cell(ws, i, 1, label, fill=WHITE, align=Alignment(horizontal="left", vertical="center"))
        write_cell(ws, i, 2, count, fill=WHITE)
        write_cell(ws, i, 3, pct / 100, fill=WHITE, number="0.0%")
    total_row = start_row + 2 + len(rows)
    total_count = sum(count for _, count, _ in rows)
    write_cell(ws, total_row, 1, "Total", fill=CREAM, font=Font(name="Calibri", bold=True), align=Alignment(horizontal="left"))
    write_cell(ws, total_row, 2, total_count, fill=CREAM, font=Font(name="Calibri", bold=True))
    write_cell(ws, total_row, 3, 1 if total_count else 0, fill=CREAM, font=Font(name="Calibri", bold=True), number="0.0%")
    return total_row + 1


def _question_values(rows: list[tuple[str, int, float]], take: int) -> list[tuple[int, float]]:
    picked = [(count, pct) for _, count, pct in rows[:take]]
    total = sum(count for count, _ in picked)
    base = total if total else 0
    return picked + [(total, 100.0 if base else 0.0)]


def write_survey_grid(ws, start_row: int, brand: str, group) -> int:
    control = arm_block(group, "control")
    experimental = arm_block(group, "experimental")
    age_labels = [QUESTION_META["S1"]["labels"][c] for c in AGE_ORDER]
    headers = (
        ["Male", "Female", "Total"]
        + age_labels
        + ["Total"]
        + ["Male", "Female", "Total"]
        + age_labels
        + ["Total"]
    )
    last_col = 21
    write_cell(ws, start_row, 1, f"Fragrance Survey · Brand {brand}", fill=HEADER, font=WHITE_FONT, align=Alignment(horizontal="left"))
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=last_col)
    write_cell(ws, start_row + 1, 1, "", fill=CREAM)
    write_cell(
        ws,
        start_row + 1,
        2,
        f"Control (1) · {control['n']} rows",
        fill=TEAL,
        font=WHITE_FONT,
    )
    ws.merge_cells(start_row=start_row + 1, start_column=2, end_row=start_row + 1, end_column=11)
    write_cell(
        ws,
        start_row + 1,
        12,
        f"Experimental (2) · {experimental['n']} rows",
        fill=RED,
        font=WHITE_FONT,
    )
    ws.merge_cells(start_row=start_row + 1, start_column=12, end_row=start_row + 1, end_column=last_col)
    write_cell(ws, start_row + 2, 1, "", fill=CREAM)
    write_cell(ws, start_row + 2, 2, "Gender", fill=TEAL, font=WHITE_FONT)
    ws.merge_cells(start_row=start_row + 2, start_column=2, end_row=start_row + 2, end_column=4)
    write_cell(ws, start_row + 2, 5, "Age", fill=TEAL, font=WHITE_FONT)
    ws.merge_cells(start_row=start_row + 2, start_column=5, end_row=start_row + 2, end_column=11)
    write_cell(ws, start_row + 2, 12, "Gender", fill=RED, font=WHITE_FONT)
    ws.merge_cells(start_row=start_row + 2, start_column=12, end_row=start_row + 2, end_column=14)
    write_cell(ws, start_row + 2, 15, "Age", fill=RED, font=WHITE_FONT)
    ws.merge_cells(start_row=start_row + 2, start_column=15, end_row=start_row + 2, end_column=last_col)
    write_cell(ws, start_row + 3, 1, "", fill=CREAM)
    bold = Font(name="Calibri", bold=True, size=10)
    for i, label in enumerate(headers, start=2):
        fill = CREAM if label != "Total" else HEADER
        font = WHITE_FONT if label == "Total" else bold
        write_cell(ws, start_row + 3, i, label, fill=fill, font=font)
    values = (
        _question_values(control["gender"], 2)
        + _question_values(control["age"], 6)
        + _question_values(experimental["gender"], 2)
        + _question_values(experimental["age"], 6)
    )
    write_cell(ws, start_row + 4, 1, "Count", fill=WHITE, font=Font(name="Calibri", bold=True), align=Alignment(horizontal="left"))
    write_cell(ws, start_row + 5, 1, "%", fill=WHITE, font=Font(name="Calibri", bold=True), align=Alignment(horizontal="left"))
    total_cols = {4, 11, 14, 21}
    for i, (count, pct) in enumerate(values, start=2):
        fill = HEADER if i in total_cols else WHITE
        font = WHITE_FONT if i in total_cols else BODY
        write_cell(ws, start_row + 4, i, count, fill=fill, font=font)
        write_cell(ws, start_row + 5, i, pct / 100, fill=fill, font=font, number="0.0%")
    return start_row + 7


def write_sheet(wb, name: str, scoped, universe_n: int, pulled_at: str):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{name} interviews"
    ws["A1"].font = TITLE
    ws.merge_cells("A1:G1")
    ws["A2"] = (
        f"{len(scoped)} {name.lower()} rows out of {universe_n} live Kobo submissions. "
        f"Pulled {pulled_at}. Complete = reached Q10 (end of interview). "
        "Incomplete = terminated on S1/S2/S3 or did not finish."
    )
    ws["A2"].font = SUB
    ws.merge_cells("A2:U2")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 32

    write_cell(ws, 4, 1, "Complete / incomplete count", fill=HEADER, font=WHITE_FONT, align=Alignment(horizontal="left"))
    ws.merge_cells("A4:C4")
    write_cell(ws, 5, 1, "Rows", fill=CREAM, font=Font(name="Calibri", bold=True))
    write_cell(ws, 5, 2, len(scoped), fill=WHITE)
    write_cell(ws, 5, 3, (len(scoped) / universe_n) if universe_n else 0, fill=WHITE, number="0.0%")

    row = 7
    row = write_frequency_table(ws, row, "Gender · all brands", freq(scoped["S2"], GENDER_ORDER, QUESTION_META["S2"]["labels"]))
    row += 1
    row = write_frequency_table(ws, row, "Age · all brands", freq(scoped["S1"], AGE_ORDER, QUESTION_META["S1"]["labels"]))
    row += 1
    if name == "Incomplete" and len(scoped):
        reasons = scoped["status_reason"].value_counts()
        reason_rows = [(label, int(count), round((int(count) / len(scoped)) * 100, 1)) for label, count in reasons.items()]
        row = write_frequency_table(ws, row, "Why the interview stopped", reason_rows)
        row += 1

    tag_map = load_tag_map()
    for brand in tag_map.get("brands") or {}:
        group = scoped[scoped["brand"] == brand] if len(scoped) else scoped.iloc[0:0]
        row = write_survey_grid(ws, row, brand, group)
        row += 1

    row += 1
    write_cell(ws, row, 1, "Respondent list", fill=HEADER, font=WHITE_FONT, align=Alignment(horizontal="left"))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    headers = ["Submitted", "Brand", "Arm", "Cell", "Age", "Gender", "Fragrance 3 months", "Status"]
    for i, label in enumerate(headers, start=1):
        write_cell(ws, row + 1, i, label, fill=CREAM, font=Font(name="Calibri", bold=True))
    listing = scoped.sort_values(["brand", "arm", "_submission_time"], na_position="last") if len(scoped) else scoped
    if not len(listing):
        write_cell(ws, row + 2, 1, "No rows in this sheet.", fill=WHITE, align=Alignment(horizontal="left"))
        ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 2, end_column=8)
    else:
        for offset, (_, rec) in enumerate(listing.iterrows()):
            submitted = rec.get("_submission_time")
            submitted = submitted.strftime("%Y-%m-%d %H:%M") if hasattr(submitted, "strftime") else ""
            values = [
                submitted,
                rec.get("brand"),
                rec.get("arm"),
                rec.get("cell"),
                QUESTION_META["S1"]["labels"].get(rec.get("S1"), "Not asked"),
                QUESTION_META["S2"]["labels"].get(rec.get("S2"), "Not asked"),
                QUESTION_META["S3"]["labels"].get(rec.get("S3"), "Not asked"),
                rec.get("status_reason"),
            ]
            for i, value in enumerate(values, start=1):
                write_cell(ws, row + 2 + offset, i, value, fill=WHITE, align=Alignment(horizontal="left"))

    widths = [28, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A4"


def export_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime("%d%m%Y")
    return f"{stamp} Fragrance Survey (Complete - Incomplete).xlsx"


def build_workbook() -> tuple[Workbook, str, int, int]:
    frame, meta = load_union(refresh=True)
    tagged, complete_mask = classify(frame)
    pulled_at = meta.get("fetched_at") or datetime.now(timezone.utc).isoformat(timespec="seconds")
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_sheet(wb, "Complete", tagged[complete_mask].copy(), len(tagged), pulled_at)
    write_sheet(wb, "Incomplete", tagged[~complete_mask].copy(), len(tagged), pulled_at)
    complete_n = int(complete_mask.sum())
    incomplete_n = int((~complete_mask).sum())
    return wb, export_filename(), complete_n, incomplete_n


def main() -> Path:
    wb, filename, complete_n, incomplete_n = build_workbook()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    wb.save(path)
    print(f"complete={complete_n} incomplete={incomplete_n}")
    print(path)
    return path


if __name__ == "__main__":
    main()
